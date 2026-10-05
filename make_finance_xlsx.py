from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DARK = "253238"
YELLOW = "E0A21A"
YELLOW_BG = "FEF6E0"
GOLD = "A67C00"
BLUE = "4A90D9"
BLUE_BG = "EBF3FD"
CARD_BG = "F7F9FB"
LINE = "E5E7EB"
MUTED = "7B8794"
WHITE = "FFFFFF"

F = "Calibri"
MONEY = '#,##0.00" ₽"'

thin = Side(style="thin", color=LINE)
bottom = Border(bottom=thin)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "Финансы"
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = YELLOW

for col, w in {"A": 30, "B": 18, "C": 20}.items():
    ws.column_dimensions[col].width = w


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def style(rng, font=None, bg=None, align=None, border=None):
    for row in ws[rng]:
        for c in row:
            if font:
                c.font = font
            if bg:
                c.fill = fill(bg)
            if align:
                c.alignment = align
            if border:
                c.border = border


# заголовок — тонкая строка
ws.merge_cells("A1:C1")
ws["A1"] = "ФИНАНСЫ"
style("A1:C1", font=Font(F, 12, bold=True, color=DARK),
      align=Alignment(horizontal="left", vertical="bottom", indent=1),
      border=Border(bottom=Side(style="thin", color=DARK)))
ws.row_dimensions[1].height = 26
ws.row_dimensions[2].height = 6

HEADERS = [("Источник", DARK, WHITE), ("В месяц", YELLOW, DARK),
           ("Сейчас лежит", BLUE, WHITE)]
for i, (h, bg, fg) in enumerate(HEADERS, start=1):
    c = ws.cell(row=3, column=i, value=h)
    c.font = Font(F, 10, bold=True, color=fg)
    c.fill = fill(bg)
    c.alignment = Alignment(horizontal="center", vertical="center")
ws.cell(row=3, column=1).alignment = Alignment(horizontal="left",
                                               vertical="center", indent=1)
ws.row_dimensions[3].height = 24

SOURCES = [
    "Постоянные доходы",
    "Фриланс",
    "Tutor AI",
    "Дивиденды",
    "Инвестиции",
    "На карте",
    "Наличка",
]

FIRST = 4
LAST = FIRST + len(SOURCES) - 1
TOTAL = LAST + 1

for i, name in enumerate(SOURCES):
    r = FIRST + i
    ws.row_dimensions[r].height = 22
    ws.cell(row=r, column=1, value=name)
    for col in range(1, 4):
        c = ws.cell(row=r, column=col)
        c.border = bottom
        c.font = Font(F, 11, color=DARK)
        if col == 1:
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        else:
            c.number_format = MONEY
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.fill = fill(YELLOW_BG if col == 2 else BLUE_BG)

ws[f"A{TOTAL}"] = "ИТОГО"
ws[f"B{TOTAL}"] = f"=SUM(B{FIRST}:B{LAST})"
ws[f"C{TOTAL}"] = f"=SUM(C{FIRST}:C{LAST})"
style(f"A{TOTAL}:C{TOTAL}", font=Font(F, 11, bold=True, color=DARK), bg=CARD_BG,
      align=Alignment(horizontal="right", vertical="center"), border=box)
ws[f"A{TOTAL}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws[f"B{TOTAL}"].number_format = MONEY
ws[f"B{TOTAL}"].font = Font(F, 12, bold=True, color=GOLD)
ws[f"C{TOTAL}"].number_format = MONEY
ws[f"C{TOTAL}"].font = Font(F, 12, bold=True, color=BLUE)
ws.row_dimensions[TOTAL].height = 28

ws.freeze_panes = "A4"

wb.calculation.fullCalcOnLoad = True
out = "/Users/ivan/Projects/tutor_helper/Финансы_2026.xlsx"
wb.save(out)
print("saved:", out)
